from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "Docker_Para_Principiantes_Paso_a_Paso.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_code_block(document: Document, code: str) -> None:
    for line in code.strip("\n").splitlines():
        paragraph = document.add_paragraph()
        run = paragraph.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(9.5)
        paragraph.paragraph_format.left_indent = Inches(0.35)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.space_before = Pt(0)


def add_bullet(document: Document, text: str) -> None:
    document.add_paragraph(text, style="List Bullet")


document = Document()

styles = document.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"].font.size = Pt(10.5)
styles["Title"].font.name = "Arial"
styles["Title"].font.size = Pt(24)
styles["Title"].font.bold = True
styles["Heading 1"].font.name = "Arial"
styles["Heading 1"].font.size = Pt(16)
styles["Heading 1"].font.bold = True
styles["Heading 2"].font.name = "Arial"
styles["Heading 2"].font.size = Pt(13)
styles["Heading 2"].font.bold = True

title = document.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run("Docker Para Principiantes: Como Aplicarlo a un Sistema Paso a Paso")
run.font.color.rgb = RGBColor(33, 33, 33)

subtitle = document.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.add_run(
    "Guia practica para alguien que nunca ha implementado Docker y quiere aprender como llevar un sistema real a contenedores."
)

document.add_heading("1. Que es Docker y por que se usa", level=1)
document.add_paragraph(
    "Docker es una herramienta que permite empaquetar una aplicacion con todo lo que necesita para correr: "
    "codigo, dependencias, configuracion y entorno. Eso hace que el sistema pueda ejecutarse de forma "
    "consistente en tu computadora, en una VPS o en otro servidor sin tener que configurar todo manualmente cada vez."
)
add_bullet(document, "Evita el problema de: en mi maquina si funciona.")
add_bullet(document, "Hace mas simple mover un sistema a un servidor.")
add_bullet(document, "Permite aislar servicios como frontend, backend y base de datos.")
add_bullet(document, "Ayuda mucho para aprender despliegue moderno.")

document.add_heading("2. Que partes de un sistema suelen ir en Docker", level=1)
document.add_paragraph(
    "Antes de escribir codigo de Docker, primero hay que entender que componentes tiene el sistema."
)
add_bullet(document, "Frontend: por ejemplo React, Vue, Next o HTML estatico.")
add_bullet(document, "Backend: por ejemplo Express, Nest, Django, Laravel o Spring.")
add_bullet(document, "Base de datos: por ejemplo MongoDB, PostgreSQL o MySQL.")
add_bullet(document, "Servicios extra: Redis, colas, workers, proxy reverso, etc.")

document.add_heading("3. Como pensar Docker antes de escribir nada", level=1)
document.add_paragraph(
    "La primera tarea no es crear un `Dockerfile`, sino dibujar mentalmente la arquitectura del sistema."
)

table = document.add_table(rows=1, cols=4)
table.style = "Table Grid"
header = table.rows[0].cells
header[0].text = "Parte"
header[1].text = "Tecnologia"
header[2].text = "Rol"
header[3].text = "Contenedor separado"
for cell in header:
    set_cell_shading(cell, "D9EAD3")

rows = [
    ("Frontend", "React/Vite", "Interfaz del usuario", "Si"),
    ("Backend", "Express", "API y logica", "Si"),
    ("Base de datos", "MongoDB", "Persistencia", "Si"),
    ("Proxy opcional", "Nginx", "Dominios/SSL", "Despues"),
]
for item in rows:
    row = table.add_row().cells
    row[0].text = item[0]
    row[1].text = item[1]
    row[2].text = item[2]
    row[3].text = item[3]

document.add_heading("4. Que se crea primero", level=1)
document.add_paragraph(
    "Una implementacion tipica con Docker se construye en este orden:"
)
steps = [
    "Se organiza la estructura del proyecto.",
    "Se verifica que la app funcione sin Docker primero.",
    "Se crea el `Dockerfile` del backend.",
    "Se crea el `Dockerfile` del frontend.",
    "Se decide si la base de datos usara imagen oficial.",
    "Se crea el archivo `docker-compose.yml` para unir todo.",
    "Se definen variables de entorno.",
    "Se prueba localmente.",
    "Se corrigen errores de puertos, rutas, dependencias o conexion.",
    "Se despliega en VPS.",
]
for step in steps:
    document.add_paragraph(step, style="List Number")

document.add_heading("5. Estructura minima de un proyecto dockerizado", level=1)
add_code_block(
    document,
    """
mi-sistema/
|-- api/
|   |-- Dockerfile
|   |-- package.json
|   `-- src/
|-- web/
|   |-- Dockerfile
|   |-- package.json
|   `-- src/
|-- docker-compose.yml
|-- .env
`-- .env.example
""",
)

document.add_heading("6. Primer archivo importante: Dockerfile", level=1)
document.add_paragraph(
    "Un `Dockerfile` es el archivo donde le dices a Docker como construir una imagen. "
    "Una imagen es como una plantilla lista para levantar un contenedor."
)

document.add_heading("6.1 Ejemplo de Dockerfile para un backend Node", level=2)
add_code_block(
    document,
    """
FROM node:22-alpine
WORKDIR /app
COPY package*.json ./
RUN npm install
COPY . .
EXPOSE 5050
CMD ["npm", "run", "dev"]
""",
)
document.add_paragraph("Explicacion linea por linea:")
add_bullet(document, "`FROM node:22-alpine`: usa una imagen oficial de Node.")
add_bullet(document, "`WORKDIR /app`: define la carpeta de trabajo dentro del contenedor.")
add_bullet(document, "`COPY package*.json ./`: copia los archivos de dependencias.")
add_bullet(document, "`RUN npm install`: instala dependencias dentro del contenedor.")
add_bullet(document, "`COPY . .`: copia el resto del proyecto.")
add_bullet(document, "`EXPOSE 5050`: documenta el puerto usado por la app.")
add_bullet(document, '`CMD ["npm", "run", "dev"]`: comando que se ejecuta al iniciar.')

document.add_heading("6.2 Ejemplo de Dockerfile para frontend React/Vite", level=2)
add_code_block(
    document,
    """
FROM node:22-alpine AS builder
WORKDIR /app
COPY package*.json ./
RUN npm install
COPY . .
RUN npm run build

FROM node:22-alpine
WORKDIR /app
RUN npm install -g serve
COPY --from=builder /app/dist ./dist
EXPOSE 4173
CMD ["serve", "-s", "dist", "-l", "4173"]
""",
)
document.add_paragraph(
    "Aqui ya aparece un concepto clave: multi-stage build. Primero se compila y luego solo se copia el "
    "resultado final para que el contenedor sea mas limpio."
)

document.add_heading("7. Segundo archivo importante: docker-compose.yml", level=1)
document.add_paragraph(
    "Si el sistema tiene varias partes, lo normal es usar `docker-compose.yml`. Este archivo conecta los "
    "servicios entre si."
)
add_code_block(
    document,
    """
services:
  api:
    build: ./api
    ports:
      - "5050:5050"
    environment:
      PORT: 5050
      MONGO_URI: mongodb://mongo:27017/appdb
    depends_on:
      - mongo

  web:
    build: ./web
    ports:
      - "4173:4173"
    depends_on:
      - api

  mongo:
    image: mongo:7
    ports:
      - "27017:27017"
    volumes:
      - ./mongodb-data:/data/db
""",
)

document.add_paragraph("Explicacion de conceptos importantes:")
add_bullet(document, "`services`: lista de contenedores del sistema.")
add_bullet(document, "`build`: indica desde donde construir una imagen propia.")
add_bullet(document, "`image`: usa una imagen ya existente.")
add_bullet(document, "`ports`: conecta puerto local con puerto interno del contenedor.")
add_bullet(document, "`environment`: define variables dentro del contenedor.")
add_bullet(document, "`depends_on`: indica dependencias entre servicios.")
add_bullet(document, "`volumes`: sirve para persistir datos.")

document.add_heading("8. Variables de entorno", level=1)
document.add_paragraph(
    "Nunca es buena practica poner datos sensibles o configuraciones cambiantes quemadas en el codigo. "
    "Por eso se usan variables de entorno."
)
add_code_block(
    document,
    """
PORT=5050
MONGO_URI=mongodb://mongo:27017/appdb
WHATSAPP_PHONE=584120000000
VITE_API_URL=http://localhost:5050/api
""",
)
add_bullet(document, "La app las lee al arrancar.")
add_bullet(document, "Cambian segun desarrollo, pruebas o produccion.")
add_bullet(document, "En VPS y Dockploy casi siempre se administran desde panel o archivo `.env`.")

document.add_heading("9. Comandos que una persona principiante debe conocer", level=1)
document.add_paragraph("Estos son los comandos mas importantes para trabajar con Docker.")

cmd_table = document.add_table(rows=1, cols=2)
cmd_table.style = "Table Grid"
cmd_header = cmd_table.rows[0].cells
cmd_header[0].text = "Comando"
cmd_header[1].text = "Para que sirve"
for cell in cmd_header:
    set_cell_shading(cell, "FCE5CD")

commands = [
    ("docker --version", "Verificar si Docker esta instalado"),
    ("docker compose up --build", "Construir y levantar todos los servicios"),
    ("docker compose up -d", "Levantar servicios en segundo plano"),
    ("docker compose down", "Apagar y eliminar contenedores del stack"),
    ("docker ps", "Ver contenedores activos"),
    ("docker images", "Ver imagenes disponibles"),
    ("docker compose logs", "Ver logs del stack"),
    ("docker compose logs api", "Ver logs de un servicio especifico"),
    ("docker exec -it <contenedor> sh", "Entrar a un contenedor"),
    ("docker compose restart", "Reiniciar servicios"),
]
for command, meaning in commands:
    row = cmd_table.add_row().cells
    row[0].text = command
    row[1].text = meaning

document.add_heading("10. Paso a paso real para dockerizar un sistema", level=1)

document.add_heading("Paso 1. Verificar que el sistema funcione sin Docker", level=2)
document.add_paragraph(
    "Antes de usar Docker, asegúrate de que el sistema ya funciona normal en tu maquina. "
    "Docker no arregla aplicaciones rotas; solo las empaqueta."
)

document.add_heading("Paso 2. Identificar servicios", level=2)
document.add_paragraph(
    "Pregunta clave: cuantas piezas tiene mi sistema. Si tienes frontend, backend y base de datos, "
    "ya sabes que probablemente habra tres servicios."
)

document.add_heading("Paso 3. Crear Dockerfile del backend", level=2)
document.add_paragraph(
    "Se escribe el archivo que indica como construir el backend."
)

document.add_heading("Paso 4. Crear Dockerfile del frontend", level=2)
document.add_paragraph(
    "Se escribe el archivo que construye el frontend y lo deja listo para servir."
)

document.add_heading("Paso 5. Definir variables", level=2)
document.add_paragraph(
    "Se crea `.env` o se preparan variables dentro de `docker-compose.yml`."
)

document.add_heading("Paso 6. Crear docker-compose.yml", level=2)
document.add_paragraph(
    "Este es el punto donde se unen todos los servicios."
)

document.add_heading("Paso 7. Construir y levantar", level=2)
add_code_block(
    document,
    """
docker compose up --build
""",
)

document.add_heading("Paso 8. Probar cada capa", level=2)
add_bullet(document, "Probar la base de datos.")
add_bullet(document, "Probar que el backend responde.")
add_bullet(document, "Probar que el frontend carga.")
add_bullet(document, "Probar que el frontend si se comunica con el backend.")

document.add_heading("Paso 9. Revisar logs si algo falla", level=2)
add_code_block(
    document,
    """
docker compose logs
docker compose logs api
docker compose logs web
docker compose logs mongo
""",
)

document.add_heading("Paso 10. Llevarlo a una VPS", level=2)
document.add_paragraph(
    "Cuando ya funciona localmente con Docker, el siguiente paso es copiar el proyecto al servidor y repetir "
    "el mismo flujo. Esa es una de las ventajas mas grandes de Docker."
)

document.add_heading("11. Errores comunes y como entenderlos", level=1)
add_bullet(document, "Puerto ocupado: otro proceso ya usa el puerto que quieres mapear.")
add_bullet(document, "No conecta a la base de datos: la URL apunta mal o el servicio no esta listo.")
add_bullet(document, "Falta una dependencia: no se instalo dentro del contenedor.")
add_bullet(document, "No encuentra archivos: el `COPY` del Dockerfile esta mal.")
add_bullet(document, "El frontend no ve el backend: la URL del API esta incorrecta.")
add_bullet(document, "Los datos se pierden: falta un volumen persistente.")

document.add_heading("12. Que codigo suele escribirse realmente", level=1)
document.add_paragraph(
    "Cuando alguien pregunta como aplicar Docker a un sistema, el codigo principal que se escribe no suele ser "
    "logica de negocio nueva, sino archivos de infraestructura:"
)
add_bullet(document, "`Dockerfile` del backend.")
add_bullet(document, "`Dockerfile` del frontend.")
add_bullet(document, "`docker-compose.yml`.")
add_bullet(document, "Archivo `.env.example`.")
add_bullet(document, "A veces scripts de arranque o seeds.")

document.add_heading("13. Flujo mental correcto para aprender Docker", level=1)
steps2 = [
    "Entender el sistema.",
    "Separar las partes.",
    "Empaquetar una parte a la vez.",
    "Conectar las partes con compose.",
    "Validar por capas.",
    "Mover el mismo stack a la VPS.",
]
for step in steps2:
    document.add_paragraph(step, style="List Number")

document.add_heading("14. Ejemplo de secuencia completa de comandos", level=1)
add_code_block(
    document,
    """
# 1. Verificar Docker
docker --version

# 2. Entrar al proyecto
cd mi-sistema

# 3. Construir y levantar
docker compose up --build

# 4. Ver contenedores
docker ps

# 5. Ver logs
docker compose logs

# 6. Apagar
docker compose down

# 7. Levantar en segundo plano
docker compose up -d
""",
)

document.add_heading("15. Recomendacion final para principiantes", level=1)
document.add_paragraph(
    "No intentes aprender Docker y una arquitectura complejisima al mismo tiempo. Empieza con un sistema "
    "pequeno pero real, como el que acabamos de construir: frontend, backend y base de datos. Cuando entiendas "
    "como se crea cada `Dockerfile`, como se conectan los servicios y como leer logs, ya tendras una base muy "
    "solida para pasar a Nginx, dominios, SSL, CI/CD y Dockploy."
)

footer = document.add_paragraph()
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_run = footer.add_run("Documento generado automaticamente para aprendizaje practico de Docker")
footer_run.italic = True

document.save(OUTPUT)
print(f"Created: {OUTPUT}")
